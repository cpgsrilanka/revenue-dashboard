#!/usr/bin/env python3
import argparse
import json
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, quote_sheetname


STANDARD_HEADERS = [
    "Agent",
    "Confirmation Date",
    "Date of Reservation",
    "Reservation No.",
    "Room Type",
    "Month",
    "Arrival",
    "Departure",
    "Guest Name",
    "Source",
    "Segment",
    "Booking Source Channel",
    "Travel Agent",
    "Nationality",
    "Rooms",
    "Room Nights",
    "Adults",
    "Children",
    "Infants",
    "Basis",
    "Marketing Code",
    "Room Charges",
    "Meal Charges",
    "Total Booking Value",
    "Payment Status",
    "Remarks",
    "Financial Year",
    "Week",
    "Pax",
]

HEADER_ALIASES = {
    "agent": "Agent",
    "revenue from santani": "Agent",
    "confirmation date": "Confirmation Date",
    "date of reservation": "Date of Reservation",
    "reservation no.": "Reservation No.",
    "reservation no": "Reservation No.",
    "reservation number": "Reservation No.",
    "room type": "Room Type",
    "villa": "Room Type",
    "month": "Month",
    "arrival": "Arrival",
    "departure": "Departure",
    "guest name": "Guest Name",
    "source": "Source",
    "segment": "Segment",
    "booking source channel": "Booking Source Channel",
    "travel agent": "Travel Agent",
    "nationality": "Nationality",
    "rooms": "Rooms",
    "room nights": "Room Nights",
    "villa nights": "Room Nights",
    "vn": "Room Nights",
    "rn": "Room Nights",
    "adults": "Adults",
    "children": "Children",
    "infants": "Infants",
    "basis": "Basis",
    "marketing code": "Marketing Code",
    "room charges": "Room Charges",
    "room charges ": "Room Charges",
    "meal charges": "Meal Charges",
    "meal charges(hb/fb/other)": "Meal Charges",
    "meal charges (hb/fb/other)": "Meal Charges",
    "total booking value": "Total Booking Value",
    "payment status": "Payment Status",
    "remarks": "Remarks",
    "financial year": "Financial Year",
    "week": "Week",
    "pax": "Pax",
}

LISTS = {
    "Segment": ["DMC", "FIT - Foreign", "FIT - Local", "OTA", "Owner", "Unassigned"],
    "Payment Status": [
        "Vouchered",
        "Deposit Paid",
        "Pending Payment",
        "Full Payment",
        "Cancelled",
        "No Show",
    ],
    "Basis": ["RO", "VO", "BB", "HB", "FB"],
}

HEADER_COMMENTS = {
    "Date of Reservation": "Booking creation date. Daily pickup is calculated from this date.",
    "Reservation No.": "Required for import. Leave blank only for planning or blocked-room template rows.",
    "Month": "Must match the arrival month, for example Jul-26.",
    "Arrival": "Guest arrival date.",
    "Departure": "Guest departure date. Must be on or after arrival unless the row is a cancellation/amendment reversal.",
    "Segment": "Use the dropdown values so dashboard segment filters stay clean.",
    "Room Nights": "Use the actual villa/room nights. Negative values are only allowed for cancellation/amendment reversals.",
    "Room Charges": "Room-only revenue. This feeds the main dashboard revenue.",
    "Meal Charges": "FNB / meal revenue. This is tracked separately from room revenue.",
    "Total Booking Value": "Should normally equal Room Charges + Meal Charges.",
    "Payment Status": "Use the dropdown values so cancellations and pending bookings are consistently handled.",
}


def normalize_header(value):
    if value is None:
        return None
    raw = str(value).replace("\n", " ").strip()
    return HEADER_ALIASES.get(raw.lower(), raw)


def collect_values(ws, col_idx, max_row):
    seen = []
    for row in range(2, max_row + 1):
        value = ws.cell(row=row, column=col_idx).value
        if value is None:
            continue
        text = str(value).strip()
        if not text or text in seen or text.startswith("="):
            continue
        seen.append(text)
    return seen[:400]


def ensure_validation_sheet(wb, ws, column_by_header, max_row):
    sheet_name = "_DashboardLists"
    if sheet_name in wb.sheetnames:
        lists_ws = wb[sheet_name]
        lists_ws.delete_rows(1, lists_ws.max_row or 1)
    else:
        lists_ws = wb.create_sheet(sheet_name)

    combined_lists = dict(LISTS)
    for header in ["Room Type", "Source", "Booking Source Channel", "Travel Agent", "Nationality", "Marketing Code"]:
        col_idx = column_by_header.get(header)
        values = collect_values(ws, col_idx, max_row) if col_idx else []
        if values:
            combined_lists[header] = values

    ranges = {}
    for list_index, (name, values) in enumerate(combined_lists.items(), start=1):
        col = get_column_letter(list_index)
        lists_ws.cell(row=1, column=list_index, value=name)
        for row_index, value in enumerate(values, start=2):
            lists_ws.cell(row=row_index, column=list_index, value=value)
        if values:
            ranges[name] = f"{quote_sheetname(sheet_name)}!${col}$2:${col}${len(values) + 1}"

    lists_ws.sheet_state = "hidden"
    lists_ws.protection.sheet = True
    lists_ws.protection.set_password("CPGRevenue2026")
    return ranges


def add_list_validation(ws, title, col_idx, formula_range, start_row, end_row):
    if not col_idx or not formula_range:
        return
    dv = DataValidation(type="list", formula1=f"={formula_range}", allow_blank=True)
    dv.error = f"Choose a valid {title} from the dropdown."
    dv.errorTitle = f"Invalid {title}"
    dv.prompt = f"Choose {title} from the list."
    dv.promptTitle = title
    ws.add_data_validation(dv)
    letter = get_column_letter(col_idx)
    dv.add(f"{letter}{start_row}:{letter}{end_row}")


def add_date_validation(ws, col_idx, start_row, end_row):
    if not col_idx:
        return
    dv = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2024,1,1)",
        formula2="DATE(2028,3,31)",
        allow_blank=True,
    )
    dv.error = "Enter a real date between 1 Jan 2024 and 31 Mar 2028."
    dv.errorTitle = "Invalid date"
    ws.add_data_validation(dv)
    letter = get_column_letter(col_idx)
    dv.add(f"{letter}{start_row}:{letter}{end_row}")


def add_custom_validation(ws, col_idx, formula, title, message, start_row, end_row):
    if not col_idx:
        return
    dv = DataValidation(type="custom", formula1=formula, allow_blank=True)
    dv.errorTitle = title
    dv.error = message
    ws.add_data_validation(dv)
    letter = get_column_letter(col_idx)
    dv.add(f"{letter}{start_row}:{letter}{end_row}")


def copy_style_from_left(ws, col_idx, max_row):
    if col_idx <= 1:
        return
    for row in range(1, min(max_row, 50) + 1):
        src = ws.cell(row=row, column=col_idx - 1)
        dst = ws.cell(row=row, column=col_idx)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format


def standardize_workbook(input_path, output_path, property_code):
    wb = load_workbook(input_path)
    if "All Bookings" not in wb.sheetnames:
        raise RuntimeError(f"{property_code}: workbook has no 'All Bookings' sheet")

    ws = wb["All Bookings"]
    report = {
        "property": property_code,
        "file": Path(input_path).name,
        "renamed_headers": [],
        "added_headers": [],
        "locked_sheets": [],
        "warnings": [],
    }

    # Some Excel Online workbooks report a huge used range because formatting
    # exists far below the real booking rows. Protect a generous entry area
    # without walking thousands of empty formatted rows.
    max_row = min(max(ws.max_row, 2000), 5000)
    max_col = max(ws.max_column, len(STANDARD_HEADERS))

    seen = set()
    for col in range(1, max_col + 1):
        original = ws.cell(row=1, column=col).value
        normalized = normalize_header(original)
        if normalized in seen and normalized != original:
            normalized = str(original).strip() if original is not None else None
        if normalized and normalized != original:
            ws.cell(row=1, column=col, value=normalized)
            report["renamed_headers"].append({"from": original, "to": normalized})
        if normalized:
            seen.add(normalized)

    for header in STANDARD_HEADERS:
        if header not in seen:
            insert_at = ws.max_column + 1
            ws.cell(row=1, column=insert_at, value=header)
            copy_style_from_left(ws, insert_at, max_row)
            report["added_headers"].append(header)
            seen.add(header)

    column_by_header = {}
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header and header not in column_by_header:
            column_by_header[str(header).strip()] = col

    missing = [header for header in STANDARD_HEADERS if header not in column_by_header]
    if missing:
        raise RuntimeError(f"{property_code}: missing required headers after normalization: {missing}")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max_row}"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="143D36")
        cell.fill = header_fill
        cell.protection = Protection(locked=True)
        header = str(cell.value).strip() if cell.value else ""
        if header in HEADER_COMMENTS:
            cell.comment = Comment(HEADER_COMMENTS[header], "Revenue Dashboard")

    ranges = ensure_validation_sheet(wb, ws, column_by_header, max_row)
    for title in [
        "Segment",
        "Payment Status",
        "Basis",
        "Room Type",
        "Source",
        "Booking Source Channel",
        "Travel Agent",
        "Nationality",
        "Marketing Code",
    ]:
        add_list_validation(ws, title, column_by_header.get(title), ranges.get(title), 2, max_row)

    for title in ["Confirmation Date", "Date of Reservation", "Arrival", "Departure"]:
        add_date_validation(ws, column_by_header.get(title), 2, max_row)

    letters = {header: get_column_letter(col) for header, col in column_by_header.items()}
    status_col = letters["Payment Status"]
    remarks_col = letters["Remarks"]
    arrival_col = letters["Arrival"]
    departure_col = letters["Departure"]
    month_col = letters["Month"]

    departure_formula = (
        f'=OR(${departure_col}2="",${arrival_col}2="",${departure_col}2>=${arrival_col}2,'
        f'LOWER(${status_col}2)="cancelled",ISNUMBER(SEARCH("cancel",${remarks_col}2)),'
        f'ISNUMBER(SEARCH("amend",${remarks_col}2)))'
    )
    add_custom_validation(
        ws,
        column_by_header["Departure"],
        departure_formula,
        "Invalid stay dates",
        "Departure must be on or after arrival, except cancellation/amendment reversal rows.",
        2,
        max_row,
    )

    month_formula = f'=OR(${month_col}2="",${arrival_col}2="",TEXT(${arrival_col}2,"mmm-yy")=${month_col}2)'
    add_custom_validation(
        ws,
        column_by_header["Month"],
        month_formula,
        "Month does not match arrival",
        "Month must match the arrival month, for example Jul-26.",
        2,
        max_row,
    )

    for header in ["Rooms", "Room Nights", "Adults", "Children", "Infants", "Pax"]:
        col = column_by_header.get(header)
        if not col:
            continue
        letter = get_column_letter(col)
        formula = (
            f'=OR({letter}2="",{letter}2>=0,LOWER(${status_col}2)="cancelled",'
            f'ISNUMBER(SEARCH("cancel",${remarks_col}2)),ISNUMBER(SEARCH("amend",${remarks_col}2)))'
        )
        add_custom_validation(
            ws,
            col,
            formula,
            f"Invalid {header}",
            f"{header} should not be negative unless this is a cancellation/amendment reversal row.",
            2,
            max_row,
        )

    for header in ["Room Charges", "Meal Charges", "Total Booking Value"]:
        col = column_by_header.get(header)
        if not col:
            continue
        letter = get_column_letter(col)
        formula = (
            f'=OR({letter}2="",{letter}2>=0,LOWER(${status_col}2)="cancelled",'
            f'ISNUMBER(SEARCH("cancel",${remarks_col}2)),ISNUMBER(SEARCH("amend",${remarks_col}2)))'
        )
        add_custom_validation(
            ws,
            col,
            formula,
            f"Invalid {header}",
            f"{header} should not be negative unless this is a cancellation/amendment reversal row.",
            2,
            max_row,
        )

    total_col = letters["Total Booking Value"]
    room_col = letters["Room Charges"]
    meal_col = letters["Meal Charges"]
    mismatch_fill = PatternFill("solid", fgColor="FFF2CC")
    ws.conditional_formatting.add(
        f"{total_col}2:{total_col}{max_row}",
        FormulaRule(
            formula=[f'=AND(${total_col}2<>"",ABS(${total_col}2-(${room_col}2+${meal_col}2))>1)'],
            fill=mismatch_fill,
        ),
    )

    bad_date_fill = PatternFill("solid", fgColor="F4CCCC")
    ws.conditional_formatting.add(
        f"{arrival_col}2:{departure_col}{max_row}",
        FormulaRule(
            formula=[
                f'=AND(${arrival_col}2<>"",${departure_col}2<>"",${departure_col}2<${arrival_col}2,'
                f'LOWER(${status_col}2)<>"cancelled")'
            ],
            fill=bad_date_fill,
        ),
    )

    for row in range(2, max_row + 1):
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.deleteRows = True
    ws.protection.set_password("CPGRevenue2026")

    for sheet in wb.worksheets:
        if sheet.title == "All Bookings" or sheet.title == "_DashboardLists":
            continue
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
        sheet.protection.sort = False
        sheet.protection.set_password("CPGRevenue2026")
        report["locked_sheets"].append(sheet.title)

    if wb.security is None:
        wb.security = WorkbookProtection()
    wb.security.lockStructure = True
    wb.security.workbookPassword = "B57B"
    wb.security.lockWindows = False

    wb.properties.modified = datetime.utcnow()
    wb.save(output_path)
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--property", required=True)
    parser.add_argument("--report", required=True)
    args = parser.parse_args()

    report = standardize_workbook(args.input, args.output, args.property)
    Path(args.report).write_text(json.dumps(report, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
